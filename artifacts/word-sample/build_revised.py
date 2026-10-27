from pathlib import Path
from docx import Document
from docx.shared import Pt
base=Path('artifacts/word-sample')
source=(base/'build_sample.py').read_text(encoding='utf-8-sig')
setup=source.split("\np('意味のかたち   タイプ別実践ガイド').runs")[0]
exec(setup)
text=(base/'revised_manuscript.md').read_text(encoding='utf-8-sig')
sections=text.strip().split('\n---\n')
for index,section in enumerate(sections):
    if index: doc.add_page_break()
    marker=p('意味のかたち   運命対話者型の自己理解と実践ガイド')
    marker.runs[0].font.size=Pt(8)
    for block in section.strip().split('\n\n'):
        if block.startswith('# '):
            title=block[2:]
            if index==0: doc.add_paragraph(title.replace('の自己理解','の\n自己理解'),style='Title')
            else: doc.add_heading(title,level=1)
        elif block.startswith('## '):
            if index==0 and block[3:].startswith('思いどおり'):
                doc.add_paragraph(block[3:],style='Subtitle')
            else:
                h(block[3:])
                if index == 6:
                    doc.paragraphs[-1].paragraph_format.space_before=Pt(6)
                    doc.paragraphs[-1].paragraph_format.space_after=Pt(6)
        elif block.startswith('?'):
            field(block[1:],2)
        else: p(block)
output=base/'運命対話者型_自己理解と実践ガイド_改訂版.docx'
doc.save(output)
print('Saved revised DOCX')
print('Japanese characters:',sum(len(x.text) for x in doc.paragraphs))
print('Sections:',len(sections))
assert not any(x in text for x in ['参考文献','参考資料','［1','［2','http'])
assert len([r for r in doc.part.rels.values() if r.reltype.endswith('/hyperlink')])==0
assert len([x for x in doc.paragraphs if x.style.name=='Title'])==1
